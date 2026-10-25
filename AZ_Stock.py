import sys
import io
import os
import json
import datetime
from typing import Optional, Dict, Any, List

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')

import time
import random
import openpyxl
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
from playwright.sync_api import sync_playwright
from concurrent.futures import ThreadPoolExecutor, as_completed
import threading
import multiprocessing
import gspread
from google.oauth2 import service_account

SCRIPT_DIR = Path(__file__).resolve().parent
EXCEL_PATH = SCRIPT_DIR / "ourmodels_pincode.xlsx"
SHEET_NAME = "Amazon"

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]

MAX_PARALLEL_CITIES = 4

print_lock = threading.Lock()
progress_counter = multiprocessing.Value('i', 0)

IN_STOCK_PHRASES = (
    "in stock",
    "left in stock",
)

OUT_OF_STOCK_PHRASES = (
    "currently unavailable",
    "temporarily out of stock",
    "out of stock",
    "we don't know when or if this item will be back in stock",
    "cannot be shipped to your selected delivery location",
    "cannot be shipped to this",
    "cannot be delivered to this address",
    "cannot be delivered to your selected location",
    "can't be delivered to your selected location",
    "can't be delivered to this address",
    "delivery is not available",
    "delivery not available",
    "not deliverable",
    "please choose a different delivery location",
    "select a different delivery location",
    "no featured offers available",
    "we're sorry, this item can't be delivered",
)

BLOCKED_PHRASES = (
    "enter the characters you see below",
    "sorry, we just need to make sure you're not a robot",
    "automated access",
    "type the characters you see in this image",
)


def log(msg):
    with print_lock:
        print(msg, flush=True)


def get_gspread_client() -> Optional[gspread.Client]:
    """Authenticates gspread client using Service Account JSON (Env var or local file)"""
    env_json = os.getenv("GDRIVE_SERVICE_ACCOUNT_JSON")
    if env_json:
        try:
            info = json.loads(env_json)
            creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
            return gspread.authorize(creds)
        except Exception as e:
            log(f"[GSheet] Error loading GDRIVE_SERVICE_ACCOUNT_JSON: {e}")

    file_path = os.getenv("GDRIVE_CREDENTIALS_FILE", "service_account.json")
    if os.path.exists(file_path):
        try:
            creds = service_account.Credentials.from_service_account_file(file_path, scopes=SCOPES)
            return gspread.authorize(creds)
        except Exception as e:
            log(f"[GSheet] Error loading credentials file {file_path}: {e}")

    return None


def classify_stock_text(text, allow_delivery_phrases=False, allow_positive=True):
    text = " ".join((text or "").split())
    lower = text.lower()
    if not lower:
        return None
    if any(phrase in lower for phrase in BLOCKED_PHRASES):
        return "blocked by Amazon"
    if any(phrase in lower for phrase in OUT_OF_STOCK_PHRASES):
        return "No"
    if allow_positive and any(phrase in lower for phrase in IN_STOCK_PHRASES):
        return "Yes"
    if allow_positive and allow_delivery_phrases and (
        "get it by" in lower
        or "get it as soon as" in lower
        or "arrives" in lower
        or "free delivery" in lower
        or "scheduled delivery" in lower
    ):
        return "Yes"
    return None


def normalize_amazon_url(url):
    raw_url = str(url).strip()
    if raw_url and "://" not in raw_url:
        raw_url = "https://" + raw_url

    parts = urlsplit(raw_url)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["th"] = "1"
    query["psc"] = "1"
    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query), ""))


def set_pincode_via_api(page, pincode):
    """Sets delivery pincode via Amazon Glow API"""
    try:
        result = page.evaluate(f"""
            async () => {{
                const resp = await fetch('/portal-migration/hz/glow/address-change', {{
                    method: 'POST',
                    headers: {{'Content-Type': 'application/x-www-form-urlencoded'}},
                    body: 'locationType=LOCATION_INPUT&zipCode={pincode}&storeContext=NoStoreName&deviceType=desktop&pageType=Gateway&actionSource=desktop-modal'
                }});
                return await resp.json();
            }}
        """)
        if result.get("successful") == 1 or result.get("isAddressUpdated") == 1:
            return True, result.get("address", {}).get("city", pincode)
        return False, ""
    except Exception:
        return False, ""


def check_stock(page, url):
    try:
        time.sleep(random.uniform(0.3, 0.8))

        page.goto(normalize_amazon_url(url), wait_until="domcontentloaded", timeout=45000)

        try:
            page.wait_for_selector(
                "#availability, #outOfStock, #add-to-cart-button, #buy-now-button, #buybox, #mir-layout-DELIVERY_BLOCK, #deliveryBlockMessage, #buybox-see-all-buying-choices",
                timeout=10000,
            )
        except:
            pass

        # 1. Check Bot Block / CAPTCHA
        title = page.title().lower()
        if any(phrase in title for phrase in ("robot check", "captcha", "automated access")):
            return "blocked by Amazon"

        try:
            body_text = page.locator("body").text_content(timeout=2000) or ""
            if any(phrase in body_text.lower() for phrase in BLOCKED_PHRASES):
                return "blocked by Amazon"
        except:
            pass

        # 2. Hard OOS: #outOfStock container visible
        try:
            oos_loc = page.locator("#outOfStock")
            if oos_loc.count() > 0 and oos_loc.first.is_visible(timeout=800):
                return "No"
        except:
            pass

        # 3. Delivery restrictions check for the current pincode
        delivery_selectors = (
            "#deliveryBlockMessage",
            "#mir-layout-DELIVERY_BLOCK",
            "#delivery-message",
            "#desktop_buybox .a-alert-error",
            "#desktop_buybox .a-alert-warning",
            "#buybox .a-alert-error",
            "#buybox .a-alert-warning",
        )
        for selector in delivery_selectors:
            try:
                loc = page.locator(selector)
                if loc.count() > 0 and loc.first.is_visible(timeout=600):
                    deliv_text = loc.first.text_content(timeout=800) or ""
                    lower_deliv = " ".join(deliv_text.split()).lower()
                    if any(phrase in lower_deliv for phrase in OUT_OF_STOCK_PHRASES):
                        return "No"
            except:
                pass

        # 4. Availability container negative check
        availability_selectors = (
            "#availability",
            "#availability_feature_div",
        )
        for selector in availability_selectors:
            try:
                locator = page.locator(selector)
                if locator.count() == 0:
                    continue
                text = locator.first.text_content(timeout=1000) or ""
                stock_class = classify_stock_text(text, allow_delivery_phrases=False, allow_positive=False)
                if stock_class == "No":
                    return "No"
            except:
                pass

        # 5. Check if only "See All Buying Options" is visible without Add to Cart
        has_atc = False
        has_bn = False
        for btn_id in ("#add-to-cart-button", "#buy-now-button"):
            try:
                btn = page.locator(btn_id)
                if btn.count() > 0 and btn.first.is_visible(timeout=800) and btn.first.is_enabled(timeout=800):
                    if btn_id == "#add-to-cart-button":
                        has_atc = True
                    else:
                        has_bn = True
            except:
                pass

        try:
            sabo_btn = page.locator("#buybox-see-all-buying-choices")
            if sabo_btn.count() > 0 and sabo_btn.first.is_visible(timeout=800) and not (has_atc or has_bn):
                return "No"
        except:
            pass

        # 6. If Add to Cart or Buy Now is present:
        if has_atc or has_bn:
            try:
                buybox_text = page.locator("#desktop_buybox, #buybox").first.text_content(timeout=1000) or ""
                lower_bb = " ".join(buybox_text.split()).lower()
                if any(phrase in lower_bb for phrase in (
                    "cannot be shipped to your selected delivery location",
                    "cannot be delivered to this address",
                    "cannot be delivered to your selected location",
                    "can't be delivered to your selected location",
                    "please choose a different delivery location",
                    "delivery is not available",
                    "not deliverable",
                )):
                    return "No"
            except:
                pass
            return "Yes"

        # 7. Positive availability fallback check
        for selector in availability_selectors:
            try:
                locator = page.locator(selector)
                if locator.count() == 0:
                    continue
                text = locator.first.text_content(timeout=1000) or ""
                stock_class = classify_stock_text(text, allow_delivery_phrases=False, allow_positive=True)
                if stock_class == "Yes":
                    return "Yes"
                elif stock_class == "No":
                    return "No"
            except:
                pass

        # 8. Check whole buybox text
        try:
            buybox_text = page.locator("#desktop_buybox, #buybox").first.text_content(timeout=1000) or ""
            lower_bb = " ".join(buybox_text.split()).lower()
            if any(p in lower_bb for p in OUT_OF_STOCK_PHRASES) or "see all buying options" in lower_bb:
                return "No"
        except:
            pass

        # 9. Page body fallback
        try:
            body_text = " ".join((page.locator("body").text_content(timeout=2000) or "").split())
            result = classify_stock_text(body_text, allow_positive=False)
            if result:
                return result
            if body_text:
                return f"unknown: {body_text[:80]}"
        except:
            pass

        return "status unknown"
    except Exception as e:
        return f"visit failed: {str(e)[:60]}"


def crawl_city(city_name, pincode, products, results_dict, progress_counter, total_tasks):
    """Crawls all products for a single city using an independent browser context"""
    city_start = time.time()
    homepage_time = 0
    pincode_time = 0
    product_times = []

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage", "--disable-blink-features=AutomationControlled"]
        )
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36",
            locale="en-IN",
            timezone_id="Asia/Kolkata",
            viewport={"width": 1366, "height": 768},
        )
        context.add_init_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        page = context.new_page()

        # Homepage load and wait for glow block
        t0 = time.time()
        try:
            page.goto("https://www.amazon.in", timeout=35000)
            page.wait_for_selector("#glow-ingress-block", state="visible", timeout=15000)
        except:
            pass
        homepage_time = time.time() - t0

        # Set pincode via Glow API
        t1 = time.time()
        ok = False
        city_detected = ""
        for _ in range(2):
            ok, city_detected = set_pincode_via_api(page, pincode)
            if ok:
                break
            time.sleep(1.0)
        pincode_time = time.time() - t1

        if not ok:
            log(f"[{city_name}] Pincode setup failed ({pincode}), skipped")
            context.close()
            browser.close()
            return

        log(f"[{city_name}] {city_detected or pincode} | home {homepage_time:.1f}s | API {pincode_time:.1f}s")

        # Check each product
        for excel_row, model, url in products:
            t2 = time.time()
            result = check_stock(page, url)
            elapsed = time.time() - t2
            product_times.append(elapsed)
            results_dict[(excel_row, city_name)] = result
            with progress_counter.get_lock():
                progress_counter.value += 1
                progress = progress_counter.value
            log(f"[{city_name}] [{progress}/{total_tasks}] {model} -> {result} ({elapsed:.1f}s)")

        context.close()
        browser.close()

    city_total = time.time() - city_start
    avg_product = sum(product_times) / len(product_times) if product_times else 0
    log(f"[{city_name}] Done: {city_total:.1f}s | avg per product: {avg_product:.1f}s")


def _sync_to_local_pricetracker(rows_data: List[List[Any]], start_row: int):
    """Backs up the updated date block into local PriceTracker (6).xlsx if sheet exists"""
    for fn in ["PriceTracker (6).xlsx", "PriceTracker.xlsx"]:
        p = SCRIPT_DIR / fn
        if p.exists():
            try:
                wb = openpyxl.load_workbook(p)
                if "Stock_availability-AZ" in wb.sheetnames:
                    ws = wb["Stock_availability-AZ"]
                    for i, r_vals in enumerate(rows_data):
                        cur_r = start_row + i
                        for c_idx, val in enumerate(r_vals, 1):
                            ws.cell(row=cur_r, column=c_idx, value=val)
                    wb.save(p)
                    log(f"[Local Sync] Backed up {len(rows_data)} rows into {fn} ('Stock_availability-AZ')")
                    break
            except Exception as e:
                log(f"[Local Sync] Note on saving to {fn}: {e}")


def sync_to_google_sheet(results_dict: Dict, products: List, city_cols: List, target_date: Optional[str] = None):
    """
    Appends or updates the 18-row daily in-stock rate block directly into Google Sheet
    worksheet 'Stock_availability-AZ' with exact formula references and formatting.
    """
    gc = get_gspread_client()
    spreadsheet_id = os.getenv("GSHEET_SPREADSHEET_ID") or os.getenv("GDRIVE_FILE_ID")

    ist_now = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=5, minutes=30)))
    date_str = target_date or os.getenv("CUSTOM_DATE") or ist_now.strftime('%Y-%m-%d')

    # Prepare Philips and Coocaa models based on template rows in ourmodels_pincode.xlsx
    philips_prods = [p for p in products if p[0] <= 13]
    coocaa_prods = [p for p in products if p[0] >= 15]

    if not gc or not spreadsheet_id:
        log("[GSheet] Google Sheets credentials or GSHEET_SPREADSHEET_ID not provided. Writing to local files only.")
        # Calculate local start_row from PriceTracker (6).xlsx
        p_file = SCRIPT_DIR / "PriceTracker (6).xlsx"
        start_row = 2544
        if p_file.exists():
            try:
                wb_tmp = openpyxl.load_workbook(p_file, read_only=True)
                if "Stock_availability-AZ" in wb_tmp.sheetnames:
                    ws_tmp = wb_tmp["Stock_availability-AZ"]
                    for r in range(ws_tmp.max_row, 0, -1):
                        if ws_tmp.cell(r, 1).value or ws_tmp.cell(r, 3).value:
                            start_row = r + 1
                            break
            except Exception:
                pass
        rows_to_insert = _build_block_rows(start_row, date_str, philips_prods, coocaa_prods, city_cols, results_dict)
        _sync_to_local_pricetracker(rows_to_insert, start_row)
        return

    try:
        sh = gc.open_by_key(spreadsheet_id)
        ws_name = "Stock_availability-AZ"
        try:
            ws = sh.worksheet(ws_name)
        except gspread.WorksheetNotFound:
            log(f"[GSheet] Worksheet '{ws_name}' not found. Creating it...")
            ws = sh.add_worksheet(title=ws_name, rows=3500, cols=26)
            h1 = ['Date', None, 'AZ_INSTOCK_RATE'] + [c[1] for c in city_cols] + ['In Stock Rate', 'Grand Total', None]
            h2 = ['Date', None, 'Pin Code'] + [c[2] for c in city_cols] + [None, None, None]
            ws.update(range_name="A1:U2", values=[h1, h2], value_input_option="USER_ENTERED")

        # Find the last row with data in Stock_availability-AZ
        col_a = ws.col_values(1)
        col_c = ws.col_values(3)
        max_len = max(len(col_a), len(col_c))
        last_row = 2
        for idx in range(max_len, 0, -1):
            va = col_a[idx - 1] if idx <= len(col_a) else ""
            vc = col_c[idx - 1] if idx <= len(col_c) else ""
            if (va and str(va).strip()) or (vc and str(vc).strip()):
                last_row = idx
                break

        log(f"[GSheet] Last populated row in '{ws_name}' is {last_row}")

        # Check if today's date block already exists within recent rows to prevent duplicates
        start_row = last_row + 1
        for check_r in range(max(3, last_row - 25), last_row + 1):
            if check_r <= len(col_a):
                existing_d = str(col_a[check_r - 1]).strip()
                if existing_d and (existing_d == date_str or date_str in existing_d):
                    start_row = check_r
                    log(f"[GSheet] Found existing block for date '{date_str}' starting at Row {start_row}. Overwriting...")
                    break

        if start_row > last_row:
            log(f"[GSheet] Appending new block for date '{date_str}' at Row {start_row}...")

        rows_to_insert = _build_block_rows(start_row, date_str, philips_prods, coocaa_prods, city_cols, results_dict)
        end_row = start_row + len(rows_to_insert) - 1

        # Ensure sheet has enough row capacity
        if end_row > ws.row_count:
            ws.add_rows(end_row - ws.row_count + 100)

        # Batch write to Google Sheet with USER_ENTERED so formulas are evaluated
        range_target = f"A{start_row}:U{end_row}"
        ws.update(range_name=range_target, values=rows_to_insert, value_input_option="USER_ENTERED")
        log(f"[GSheet] Successfully wrote {len(rows_to_insert)} rows to '{ws_name}' at range {range_target}")

        # Format rows: center alignment, percent formats, and borders
        try:
            percent_format = {"type": "PERCENT", "pattern": "0.00%"}
            format_requests = [
                {
                    "repeatCell": {
                        "range": {
                            "sheetId": ws.id,
                            "startRowIndex": start_row - 1,
                            "endRowIndex": end_row,
                            "startColumnIndex": 0,
                            "endColumnIndex": 21
                        },
                        "cell": {
                            "userEnteredFormat": {
                                "horizontalAlignment": "CENTER",
                                "textFormat": {"fontSize": 10, "fontFamily": "Calibri"}
                            }
                        },
                        "fields": "userEnteredFormat(horizontalAlignment,textFormat)"
                    }
                },
                {
                    "repeatCell": {
                        "range": {
                            "sheetId": ws.id,
                            "startRowIndex": start_row - 1,
                            "endRowIndex": end_row,
                            "startColumnIndex": 18,
                            "endColumnIndex": 19
                        },
                        "cell": {
                            "userEnteredFormat": {
                                "numberFormat": percent_format
                            }
                        },
                        "fields": "userEnteredFormat.numberFormat"
                    }
                },
                {
                    "repeatCell": {
                        "range": {
                            "sheetId": ws.id,
                            "startRowIndex": start_row - 1,
                            "endRowIndex": end_row,
                            "startColumnIndex": 20,
                            "endColumnIndex": 21
                        },
                        "cell": {
                            "userEnteredFormat": {
                                "numberFormat": percent_format
                            }
                        },
                        "fields": "userEnteredFormat.numberFormat"
                    }
                },
                {
                    "updateBorders": {
                        "range": {
                            "sheetId": ws.id,
                            "startRowIndex": start_row - 1,
                            "endRowIndex": end_row,
                            "startColumnIndex": 0,
                            "endColumnIndex": 21
                        },
                        "top": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.8, "blue": 0.8}},
                        "bottom": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.8, "blue": 0.8}},
                        "left": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.8, "blue": 0.8}},
                        "right": {"style": "SOLID", "width": 1, "color": {"red": 0.8, "green": 0.8, "blue": 0.8}},
                        "innerHorizontal": {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.85, "blue": 0.85}},
                        "innerVertical": {"style": "SOLID", "width": 1, "color": {"red": 0.85, "green": 0.85, "blue": 0.85}}
                    }
                }
            ]
            sh.batch_update({"requests": format_requests})
            log("[GSheet] Formatting, percentage number formats, and borders applied successfully.")
        except Exception as fe:
            log(f"[GSheet] Note on formatting: {fe}")

        # Local PriceTracker backup
        _sync_to_local_pricetracker(rows_to_insert, start_row)

    except Exception as e:
        log(f"[GSheet] Error updating Google Sheets: {e}")


def _build_block_rows(start_row: int, date_str: str, philips_prods: List, coocaa_prods: List, city_cols: List, results_dict: Dict) -> List[List[Any]]:
    """Builds the 18 rows structure matching Stock_availability-AZ template"""
    rows = []

    # 1. Philips models (11 models)
    for idx, (ex_row, model, url) in enumerate(philips_prods):
        cur_r = start_row + idx
        city_vals = []
        for _, c_name, _ in city_cols:
            raw_res = results_dict.get((ex_row, c_name), "No")
            city_vals.append("Yes" if str(raw_res).strip().lower() == "yes" else "No")

        rate_formula = "-%" if "85PQT8101" in model else f'=COUNTIF(D{cur_r}:R{cur_r},"Yes")/15'
        rows.append([date_str, url, model] + city_vals + [rate_formula, "", ""])

    # 2. Philips Summary row (Row 12 of the block)
    p_start1 = start_row
    p_end1 = start_row + 6
    p_start2 = start_row + 8
    p_end2 = start_row + 10
    philips_summary_formula = f'=(COUNTIF(D{p_start1}:R{p_end1},"Yes")+COUNTIF(D{p_start2}:R{p_end2},"Yes"))/150'
    summary_row = [date_str, "", ""] + [""] * 13 + ["AZ - Philips", "", philips_summary_formula, "", ""]
    rows.append(summary_row)

    # 3. Coocaa models (6 models)
    c_block_start = start_row + len(rows)
    for idx, (ex_row, model, url) in enumerate(coocaa_prods):
        cur_r = c_block_start + idx
        city_vals = []
        for _, c_name, _ in city_cols:
            raw_res = results_dict.get((ex_row, c_name), "No")
            city_vals.append("Yes" if str(raw_res).strip().lower() == "yes" else "No")

        rate_formula = f'=COUNTIF(D{cur_r}:R{cur_r},"Yes")/15'
        if idx == 1:  # 55Y74 Pro holds AZ TOTAL
            c_start = c_block_start
            c_end = c_block_start + len(coocaa_prods) - 1
            az_total_formula = f'=(COUNTIF(D{c_start}:R{c_end},"Yes")+COUNTIF(D{p_start1}:R{p_end1},"Yes")+COUNTIF(D{p_start2}:R{p_end2},"Yes"))/240'
            row_data = [date_str, url, model] + city_vals + [rate_formula, "AZ TOTAL", az_total_formula]
        else:
            row_data = [date_str, url, model] + city_vals + [rate_formula, "", ""]
        rows.append(row_data)

    return rows


def main():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(f"Template workbook not found: {EXCEL_PATH}")

    wb = openpyxl.load_workbook(EXCEL_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' not found in {EXCEL_PATH}")
    ws = wb[SHEET_NAME]

    headers = [cell.value for cell in ws[1]]
    pincodes = [cell.value for cell in ws[2]]

    # 15 city columns (from 3rd column onwards)
    city_cols = []
    for col_idx in range(2, len(headers)):
        if col_idx < len(pincodes):
            city_name = headers[col_idx]
            pincode = pincodes[col_idx]
            if city_name and pincode:
                city_cols.append((col_idx + 1, str(city_name).strip(), str(int(pincode))))

    # Product list (from row 3)
    products = []
    for excel_row, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        model = row[0]
        url = row[1]
        if model and url:
            products.append((excel_row, model, url))

    # Optional test limit from environment
    test_limit_str = os.getenv("TEST_LIMIT", "")
    if test_limit_str.isdigit():
        t_limit = int(test_limit_str)
        if 0 < t_limit < len(products):
            print(f"Applying TEST_LIMIT: testing only first {t_limit} products.")
            products = products[:t_limit]

    total_products = len(products)
    total_cities = len(city_cols)
    total_tasks = total_products * total_cities

    print(f"Products: {total_products}")
    print(f"Cities: {total_cities} x {MAX_PARALLEL_CITIES} parallel")
    print(f"Total checks: {total_tasks}")
    for _, city, pincode in city_cols:
        print(f"   - {city}: {pincode}")
    print()

    results_dict = {}
    global progress_counter
    progress_counter.value = 0
    start_time = time.time()

    with ThreadPoolExecutor(max_workers=MAX_PARALLEL_CITIES) as executor:
        futures = {
            executor.submit(crawl_city, city_name, pincode, products, results_dict, progress_counter, total_tasks): city_name
            for _, city_name, pincode in city_cols
        }

        for future in as_completed(futures):
            city_name = futures[future]
            try:
                future.result()
            except Exception as e:
                log(f"[{city_name}] Error: {e}")

    total_time = time.time() - start_time

    print(f"\n{'='*60}")
    print(f"All done!")
    print(f"  Total time: {total_time:.1f}s ({total_time/60:.1f} min)")
    print(f"  Products: {total_products}")
    print(f"  Cities: {total_cities}")
    print(f"  Total checks: {total_tasks}")
    print(f"  Average speed: {total_time/total_tasks:.1f}s/check")
    print(f"{'='*60}")

    # Write back to ourmodels_pincode.xlsx
    for row_idx in range(3, ws.max_row + 1):
        model = ws.cell(row=row_idx, column=1).value
        url = ws.cell(row=row_idx, column=2).value
        if not model and not url:
            for col_idx, _, _ in city_cols:
                ws.cell(row=row_idx, column=col_idx).value = None

    for (excel_row, city_name), result in results_dict.items():
        for col_idx, c_name, _ in city_cols:
            if c_name == city_name:
                ws.cell(row=excel_row, column=col_idx, value=result)
                break

    for attempt in range(3):
        try:
            wb.save(EXCEL_PATH)
            print(f"\nResults written to template: {EXCEL_PATH}")
            break
        except PermissionError:
            if attempt < 2:
                print(f"Excel file is open, retrying in {attempt+2}s...")
                time.sleep(attempt + 2)
            else:
                backup_path = EXCEL_PATH.with_name(EXCEL_PATH.stem + "_backup.xlsx")
                wb.save(backup_path)
                print(f"Original file is open, results saved to: {backup_path}")

    # Sync to Google Sheet worksheet 'Stock_availability-AZ'
    sync_to_google_sheet(results_dict, products, city_cols)


if __name__ == "__main__":
    main()
