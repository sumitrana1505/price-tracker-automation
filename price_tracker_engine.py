"""
Price Tracker Brain - Daily Execution & Cloud-Enabled Engine
=============================================================
Reads competitor links from 'Competitor price monitoring',
scrapes live Price, Bank Offers, Coupons, Ratings, and Ratings Count,
and writes them into today's 6-column block in the workbook.
Supports local execution and cloud execution via Google Drive sync.
"""

import argparse
import copy
import logging
import os
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Optional, Dict, Any, List

# Windows console UTF-8 output configuration
if sys.platform == "win32":
    os.environ.setdefault("PYTHONIOENCODING", "utf-8")
    if sys.stdout and hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if sys.stderr and hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import (
    DEFAULT_WORKBOOK,
    FALLBACK_WORKBOOK,
    SHEET_NAME,
    FIRST_DATA_ROW,
    LAST_DATA_ROW,
    COL_CHANNEL,
    COL_MODEL,
    COL_LINK,
    BLOCK_SIZE,
    SUB_HEADERS,
    format_today_header,
)
from scrapers import FlipkartScraper, AmazonScraper, ProductScrapedData
from gdrive_sync import GoogleDriveSync

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
    handlers=[
        logging.StreamHandler(sys.stdout),
    ]
)
logger = logging.getLogger("PriceTrackerBrain")


class PriceTrackerBrain:
    def __init__(self, file_path: str, date_header: Optional[str] = None):
        self.file_path = file_path
        self.date_header = date_header or format_today_header()
        self.wb = None
        self.ws = None
        self.fk_scraper = FlipkartScraper()
        self.amz_scraper = AmazonScraper()

    def load_workbook(self):
        """Loads target Excel workbook safely"""
        if not os.path.exists(self.file_path):
            raise FileNotFoundError(f"Workbook not found at: {self.file_path}")
        
        logger.info(f"Loading workbook: {self.file_path}")
        self.wb = openpyxl.load_workbook(self.file_path)
        if SHEET_NAME not in self.wb.sheetnames:
            raise ValueError(f"Sheet '{SHEET_NAME}' not found in {self.wb.sheetnames}")
        self.ws = self.wb[SHEET_NAME]
        logger.info(f"Loaded sheet '{SHEET_NAME}' ({self.ws.max_row} rows, {self.ws.max_column} cols)")

    def find_or_create_date_columns(self) -> int:
        """
        Locates existing 6-column block for today's date, or appends a new block.
        Returns the 1-based start column index.
        """
        ws = self.ws
        max_col = ws.max_column

        # Search existing Row 1 headers for target date
        for c in range(21, max_col + 1):
            val = ws.cell(1, c).value
            if val is not None and str(val).strip().lower() == self.date_header.strip().lower():
                logger.info(f"Found existing column block for '{self.date_header}' at Col {c} ({get_column_letter(c)})")
                return c

        # Block does not exist; determine last used column in row 2
        last_used_col = 20
        for c in range(21, max_col + 30):
            val = ws.cell(2, c).value
            if val is not None and str(val).strip() != "":
                last_used_col = c

        new_start_col = last_used_col + 1
        end_col = new_start_col + BLOCK_SIZE - 1
        logger.info(f"Creating NEW column block for '{self.date_header}' at Col {new_start_col}-{end_col} ({get_column_letter(new_start_col)}-{get_column_letter(end_col)})")

        thin_side = Side(border_style="thin", color="000000")
        all_borders = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        header_font = Font(name="Arial", size=8, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3D85C6", end_color="3D85C6", fill_type="solid")
        prev_date_cell = ws.cell(1, new_start_col - BLOCK_SIZE)
        if prev_date_cell.fill and prev_date_cell.fill.fill_type:
            header_fill = copy.copy(prev_date_cell.fill)

        # 1. Header Row 1: Date label (White text, blue fill, Arial 8, bold, Merge and Center with All Borders)
        for c in range(new_start_col, end_col + 1):
            c_cell = ws.cell(1, c)
            c_cell.border = all_borders
            c_cell.fill = header_fill
            c_cell.font = header_font

        date_cell = ws.cell(1, new_start_col, self.date_header)
        date_cell.alignment = Alignment(horizontal="center", vertical="center")
        date_cell.font = header_font
        ws.merge_cells(start_row=1, start_column=new_start_col, end_row=1, end_column=end_col)

        # 2. Header Row 2: Sub-headers (White text, blue fill, Arial 8, bold, Centered with All Borders)
        for i, sub in enumerate(SUB_HEADERS):
            c_idx = new_start_col + i
            sub_cell = ws.cell(2, c_idx, sub)
            sub_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sub_cell.font = header_font
            sub_cell.border = all_borders
            sub_cell.fill = header_fill

        # 3. Format All Borders, Black Font, Arial 8, Center Alignment for all data rows (Row 3 to 249)
        data_font = Font(name="Arial", size=8, bold=False, color="000000")
        data_align = Alignment(horizontal="center", vertical="center")
        for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            for c in range(new_start_col, end_col + 1):
                data_c = ws.cell(r, c)
                data_c.border = all_borders
                data_c.alignment = data_align
                data_c.font = data_font

        return new_start_col

    def execute(self, test_limit: Optional[int] = None, skip_existing: bool = False):
        """
        Runs the data fetching and Excel updating process.
        test_limit: If set, only processes first N rows with valid links.
        skip_existing: If True, skips rows that already have Price filled.
        """
        self.load_workbook()
        start_col = self.find_or_create_date_columns()
        
        col_price = start_col
        col_bank = start_col + 1
        col_coupon = start_col + 2
        col_nep = start_col + 3
        col_rating = start_col + 4
        col_rc = start_col + 5

        col_price_letter = get_column_letter(col_price)
        col_bank_letter = get_column_letter(col_bank)
        col_coupon_letter = get_column_letter(col_coupon)

        ws = self.ws
        rows_to_process = []

        for r in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
            channel = str(ws.cell(r, COL_CHANNEL).value or "").strip()
            model = str(ws.cell(r, COL_MODEL).value or "").strip()
            link = str(ws.cell(r, COL_LINK).value or "").strip()

            if not channel and not model and not link:
                continue

            rows_to_process.append({
                "row": r,
                "channel": channel,
                "model": model,
                "link": link
            })

        logger.info(f"Total tracker rows to monitor: {len(rows_to_process)}")
        
        processed_count = 0
        success_count = 0
        fail_count = 0
        t0 = time.time()

        for item in rows_to_process:
            r = item["row"]
            ch = item["channel"].lower()
            model = item["model"]
            link = item["link"]

            # Check if row already has price and skip_existing is True
            existing_price = ws.cell(r, col_price).value
            if skip_existing and existing_price is not None:
                continue

            # Always write NEP formula
            nep_formula = f"=IFERROR({col_price_letter}{r}-{col_bank_letter}{r}-{col_coupon_letter}{r},0)"
            ws.cell(r, col_nep, nep_formula)

            # Check if link is empty / placeholder
            if not link or not link.startswith("http"):
                ws.cell(r, col_price, None)
                ws.cell(r, col_bank, None)
                ws.cell(r, col_coupon, 0)
                ws.cell(r, col_rating, None)
                ws.cell(r, col_rc, None)
                continue

            if test_limit and processed_count >= test_limit:
                logger.info(f"Reached test limit ({test_limit} products). Stopping test run.")
                break

            processed_count += 1
            logger.info(f"[{processed_count}] Row {r} | {ch.capitalize()} | Model: {model}")

            scraped_data = None
            try:
                if "flipkart" in ch or "flipkart" in link:
                    scraped_data = self.fk_scraper.scrape(link, model=model)
                elif "amazon" in ch or "amazon" in link:
                    scraped_data = self.amz_scraper.scrape(link, model=model)
                else:
                    logger.warning(f"  Unknown channel '{ch}' for row {r}")
            except Exception as e:
                logger.error(f"  Scrape exception on row {r}: {e}")

            if scraped_data and isinstance(scraped_data.price, int):
                # Write Scraped Data
                ws.cell(r, col_price, scraped_data.price)
                ws.cell(r, col_bank, scraped_data.bank_offer)
                ws.cell(r, col_coupon, scraped_data.coupon or 0)
                ws.cell(r, col_rating, scraped_data.rating)
                ws.cell(r, col_rc, scraped_data.rating_count)

                nep_calc = scraped_data.price - (scraped_data.bank_offer or 0) - (scraped_data.coupon or 0)
                logger.info(
                    f"  [SUCCESS] Price: Rs.{scraped_data.price:,} | "
                    f"Bank Offer: Rs.{scraped_data.bank_offer or 0:,} | "
                    f"Coupon: Rs.{scraped_data.coupon or 0:,} | "
                    f"NEP: Rs.{nep_calc:,} | "
                    f"Rating: {scraped_data.rating} ({scraped_data.rating_count or 0})"
                )
                success_count += 1
            elif scraped_data and (scraped_data.price == "OOS" or not scraped_data.in_stock):
                logger.info(f"  [OOS] Marked as OOS across all pincodes | Rating: {scraped_data.rating} ({scraped_data.rating_count or 0})")
                ws.cell(r, col_price, "OOS")
                ws.cell(r, col_bank, None)
                ws.cell(r, col_coupon, 0)
                ws.cell(r, col_rating, scraped_data.rating)
                ws.cell(r, col_rc, scraped_data.rating_count)
                fail_count += 1
            else:
                err = scraped_data.error_msg if scraped_data else "Failed to fetch"
                logger.warning(f"  [OOS / ERROR] {err}")
                ws.cell(r, col_price, "OOS")
                ws.cell(r, col_bank, None)
                ws.cell(r, col_coupon, 0)
                if scraped_data:
                    ws.cell(r, col_rating, scraped_data.rating)
                    ws.cell(r, col_rc, scraped_data.rating_count)
                fail_count += 1

            # Auto-save every 5 products
            if processed_count % 5 == 0:
                self.save_workbook()

        # Final save
        self.save_workbook()
        elapsed = time.time() - t0
        logger.info("=" * 60)
        logger.info("  CRAWL COMPLETE")
        logger.info(f"  Total Processed: {processed_count}")
        logger.info(f"  Successful: {success_count} | OOS/Failed: {fail_count}")
        logger.info(f"  Elapsed: {elapsed:.1f}s ({elapsed/60:.1f} mins)")
        logger.info("=" * 60)

    def save_workbook(self):
        """Safe workbook save handling Windows file locks"""
        for attempt in range(1, 4):
            try:
                self.wb.save(self.file_path)
                logger.info(f"Auto-saved progress to '{self.file_path}'")
                return
            except PermissionError:
                logger.warning(f"File locked by Excel! Attempt {attempt}/3...")
                time.sleep(2)
            except Exception as e:
                logger.error(f"Save failed: {e}")
                break


def resolve_default_file() -> str:
    """Finds available workbook in workspace"""
    for candidate in [DEFAULT_WORKBOOK, FALLBACK_WORKBOOK]:
        if os.path.exists(candidate):
            return candidate
    for f in os.listdir("."):
        if f.lower().startswith("pricetracker") and f.endswith(".xlsx"):
            return f
    return DEFAULT_WORKBOOK


def main():
    parser = argparse.ArgumentParser(description="Price Tracker Automation Brain")
    parser.add_argument("--file", default=None, help="Path to PriceTracker Excel file")
    parser.add_argument("--date", default=None, help="Custom date header (e.g. '16th - Oct')")
    parser.add_argument("--test", type=int, default=None, help="Run test on first N products")
    parser.add_argument("--skip-existing", action="store_true", help="Skip rows that already have price data")
    parser.add_argument("--sync-gdrive", action="store_true", help="Download from and upload back to Google Drive")
    parser.add_argument("--gsheet", default=None, help="Direct Google Sheet Spreadsheet ID")
    args = parser.parse_args()

    # Read from CLI flags or environment variables
    spreadsheet_id = args.gsheet or os.getenv("GSHEET_SPREADSHEET_ID")
    target_date = args.date or os.getenv("CUSTOM_DATE") or format_today_header()
    # Strip any accidental wrapping quotes
    target_date = target_date.strip().strip("'\"") if target_date else format_today_header()
    
    test_limit = args.test
    if test_limit is None and os.getenv("TEST_LIMIT"):
        try:
            test_limit = int(os.getenv("TEST_LIMIT").strip())
        except ValueError:
            test_limit = None

    if spreadsheet_id:
        from gsheet_engine import GoogleSheetPriceTrackerBrain
        logger.info("=" * 60)
        logger.info("  DIRECT GOOGLE SHEETS MODE")
        logger.info(f"  Spreadsheet ID: {spreadsheet_id}")
        logger.info(f"  Date Header: {target_date}")
        if test_limit:
            logger.info(f"  Test Limit: {test_limit} products")
        logger.info("=" * 60)
        gsheet_brain = GoogleSheetPriceTrackerBrain(spreadsheet_id=spreadsheet_id, date_header=target_date)
        gsheet_brain.run(test_limit=test_limit)
        return

    target_file = args.file or resolve_default_file()

    logger.info("=" * 60)
    logger.info("  PRICE TRACKER AUTOMATION BRAIN")
    logger.info(f"  Target File: {target_file}")
    logger.info(f"  Date Header: {args.date or format_today_header()}")
    logger.info(f"  Google Drive Sync: {'Enabled' if args.sync_gdrive else 'Disabled'}")
    logger.info("=" * 60)

    gdrive = GoogleDriveSync()
    if args.sync_gdrive and gdrive.is_configured():
        logger.info("Syncing latest version from Google Drive before crawl...")
        gdrive.download_file(target_file)

    brain = PriceTrackerBrain(file_path=target_file, date_header=args.date)
    brain.execute(test_limit=args.test, skip_existing=args.skip_existing)

    if args.sync_gdrive and gdrive.is_configured():
        logger.info("Uploading updated workbook to Google Drive...")
        gdrive.upload_file(target_file)


if __name__ == "__main__":
    main()
